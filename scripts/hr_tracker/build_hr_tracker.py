#!/usr/bin/env python3
"""
Build the HR task tracker outputs from hr_tasks_data.py:
  - docs/hr/HR_Task_Tracker_Oct2026.xlsx   (uploaded to Google Drive as the living Google Sheet)
  - docs/hr/HR_Task_Tracker_Oct2026.md     (repo copy of the tracker)
  - docs/hr/HR_Roles_Responsibilities_Additions_Oct2026.md
  - <out>/clickup_tasks.json               (payload for ClickUp create_many)
  - <out>/rr_additions.html                (for Google Doc upload)

Usage: python3 build_hr_tracker.py [output_dir]
"""
import json
import os
import sys
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, os.path.dirname(__file__))
from hr_tasks_data import DEADLINE, OWNERS, RR_ADDITIONS, SNAPSHOT_DATE, TASKS  # noqa: E402

REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
DOCS = os.path.join(REPO, "docs", "hr")
OUT = sys.argv[1] if len(sys.argv) > 1 else DOCS
os.makedirs(DOCS, exist_ok=True)
os.makedirs(OUT, exist_ok=True)

OPEN_STATUSES = {"PENDING", "PARTIAL", "VERIFY"}
STATUS_FILL = {
    "COMPLETED": "C6EFCE",
    "PENDING": "FFC7CE",
    "PARTIAL": "FFEB9C",
    "VERIFY": "DDEBF7",
    "RECURRING": "E4DFEC",
}
PRIORITY_ORDER = {"urgent": 0, "high": 1, "normal": 2, "low": 3}

HEADERS = [
    "ID", "Category", "Task / Instruction", "Owner", "Status", "Priority",
    "Source (chat & date)", "Evidence / last update", "Next action", "Deadline",
    "ClickUp task", "Follow-up date", "Latest remark (update here)",
]


def row_for(t):
    deadline = DEADLINE if t["status"] in OPEN_STATUSES else ("Monthly / ongoing" if t["status"] == "RECURRING" else "")
    return [
        t["id"], t["category"], t["task"], OWNERS[t["owner"]], t["status"], t["priority"].upper(),
        t["source"], t["evidence"], t["next_action"], deadline,
        "Yes" if t["clickup"] else "", "", "",
    ]


def style_sheet(ws, n_rows, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2, max_row=n_rows + 1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"


def add_table(ws, name, n_rows, n_cols):
    ref = f"A1:{get_column_letter(n_cols)}{n_rows + 1}"
    tbl = Table(displayName=name, ref=ref)
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tbl)


def colour_status(ws, status_col, n_rows):
    for r in range(2, n_rows + 2):
        c = ws.cell(row=r, column=status_col)
        fill = STATUS_FILL.get(str(c.value))
        if fill:
            c.fill = PatternFill("solid", fgColor=fill)
            c.font = Font(bold=True)


def build_xlsx(path):
    wb = Workbook()

    # README
    ws = wb.active
    ws.title = "Read Me"
    readme = [
        ["HR Department - Task Tracker (Living Document)"],
        [f"Snapshot built from WhatsApp exports on {SNAPSHOT_DATE}. Deadline for all open items: {DEADLINE} (10 Nov 2026)."],
        [""],
        ["How to use"],
        ["1. 'Master Task Log' = every instruction found in the HR WhatsApp chats (group + 1:1 with Rajendra Prasad + 1:1 with Gnaneshwar), with status and evidence."],
        ["2. 'Pending Action Items' = only the open items (PENDING / PARTIAL / VERIFY) sorted by priority - use this for weekly follow-up."],
        ["3. Update the 'Status' and 'Latest remark' columns in Master Task Log; the Pending tab is a filtered copy, so re-copy or filter the master when you review."],
        ["4. Add new instructions as new rows with the next ID in the category (A = Performance, B = ClickUp/Daily reports, C = Recruitment, D = Onboarding/Trusts, E = Payroll, F = Engagement, G = Policies, H = Employee cases)."],
        ["5. Log every change in 'Change Log' (date, who, what changed)."],
        ["6. ClickUp: open items are also created as tasks assigned to Gnaneshwar (Rajendra Prasad does not use ClickUp) in HR Department > 'MKCD Follow-up Action Items (WhatsApp) - Due 10 Nov 2026'."],
        [""],
        ["Status legend"],
        ["COMPLETED - evidence of completion in chat"],
        ["PENDING - instruction given, no evidence of completion"],
        ["PARTIAL - started / partly delivered, remainder open"],
        ["VERIFY - probably done but no written confirmation; HR to confirm"],
        ["RECURRING - standing process (daily / monthly); shown for R&R and routine checks"],
        [""],
        ["Owners"],
        ["RP = Rajendra Prasad (HR Manager); GN = Gnaneshwar (HR Executive); BOTH = joint; MKCD = your own decision pending"],
    ]
    for r in readme:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=14)
    ws["A4"].font = Font(bold=True)
    ws["A12"].font = Font(bold=True)
    ws["A19"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 140
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    widths = [6, 22, 60, 24, 12, 10, 30, 50, 50, 14, 10, 14, 40]

    # Master
    ws = wb.create_sheet("Master Task Log")
    ws.append(HEADERS)
    for t in TASKS:
        ws.append(row_for(t))
    style_sheet(ws, len(TASKS), widths)
    colour_status(ws, 5, len(TASKS))
    add_table(ws, "MasterTasks", len(TASKS), len(HEADERS))

    # Pending
    pending = sorted(
        [t for t in TASKS if t["status"] in OPEN_STATUSES],
        key=lambda t: (PRIORITY_ORDER[t["priority"]], t["id"]),
    )
    ws = wb.create_sheet("Pending Action Items")
    ws.append(HEADERS)
    for t in pending:
        ws.append(row_for(t))
    style_sheet(ws, len(pending), widths)
    colour_status(ws, 5, len(pending))
    add_table(ws, "PendingTasks", len(pending), len(HEADERS))

    # Recurring
    recurring = [t for t in TASKS if t["status"] == "RECURRING"]
    ws = wb.create_sheet("Recurring Processes")
    ws.append(HEADERS)
    for t in recurring:
        ws.append(row_for(t))
    style_sheet(ws, len(recurring), widths)
    colour_status(ws, 5, len(recurring))
    add_table(ws, "RecurringTasks", len(recurring), len(HEADERS))

    # R&R additions
    ws = wb.create_sheet("R&R Additions")
    ws.append(["Role", "Responsibility", "Details / Standard", "Frequency"])
    freq_guess = {
        "Daily": ["Daily", "daily", "closing", "attendance report"],
        "Weekly": ["Weekly", "weekly"],
        "Monthly": ["Monthly", "monthly", "payroll", "Pay sheets", "15th"],
    }
    for role, items in RR_ADDITIONS.items():
        for title, detail in items:
            freq = "As required"
            for f, keys in freq_guess.items():
                if any(k in title or k in detail for k in keys):
                    freq = f
                    break
            ws.append([role, title, detail, freq])
    n = sum(len(v) for v in RR_ADDITIONS.values())
    style_sheet(ws, n, [40, 32, 100, 14])
    add_table(ws, "RRAdditions", n, 4)

    # Change log
    ws = wb.create_sheet("Change Log")
    ws.append(["Date", "Updated by", "Task ID", "Change / remark"])
    ws.append([SNAPSHOT_DATE, "Claude (for MKCD)", "ALL", "Initial tracker built from WhatsApp exports (group chat Aug 2025 - Oct 2026, 1:1 chats with Rajendra Prasad and Gnaneshwar)."])
    style_sheet(ws, 1, [14, 20, 10, 100])

    wb.save(path)
    return pending, recurring


def md_table(rows, cols):
    def esc(v):
        return str(v).replace("|", "\\|").replace("\n", " ")
    out = ["| " + " | ".join(cols) + " |", "|" + "---|" * len(cols)]
    for r in rows:
        out.append("| " + " | ".join(esc(x) for x in r) + " |")
    return "\n".join(out)


def build_markdown(pending, recurring, sheet_url=None, clickup_url=None, doc_url=None):
    counts = {}
    for t in TASKS:
        counts[t["status"]] = counts.get(t["status"], 0) + 1
    lines = [
        "# HR Department - Task Tracker (living document)",
        "",
        f"Built from WhatsApp exports on {SNAPSHOT_DATE}. Deadline for all open items: **10 November 2026**.",
        "",
        "Sources: Group#87_HR Tasks (Aug 2025 - Oct 2026), 1:1 chat with Rajendra Prasad (HR Manager), 1:1 chat with Gnaneshwar (HR Executive).",
        "",
    ]
    if sheet_url or clickup_url or doc_url:
        lines.append("## Live copies")
        if sheet_url:
            lines.append(f"- Google Sheet (update here): {sheet_url}")
        if clickup_url:
            lines.append(f"- ClickUp list (tasks assigned to Gnaneshwar): {clickup_url}")
        if doc_url:
            lines.append(f"- Roles & Responsibilities additions (Google Doc): {doc_url}")
        lines.append("")
    lines += [
        "## Summary",
        "",
        md_table([[k, v] for k, v in sorted(counts.items())], ["Status", "Count"]),
        "",
        "## Pending action items (follow up with HR)",
        "",
        md_table(
            [[t["id"], t["priority"].upper(), t["status"], OWNERS[t["owner"]], t["task"], t["next_action"]] for t in pending],
            ["ID", "Priority", "Status", "Owner", "Task", "Next action"],
        ),
        "",
        "## Recurring processes (standing expectations)",
        "",
        md_table([[t["id"], OWNERS[t["owner"]], t["task"]] for t in recurring], ["ID", "Owner", "Process"]),
        "",
        "## Master task log",
        "",
        md_table(
            [[t["id"], t["category"], t["task"], OWNERS[t["owner"]], t["status"], t["source"], t["evidence"]] for t in TASKS],
            ["ID", "Category", "Task", "Owner", "Status", "Source", "Evidence / last update"],
        ),
        "",
    ]
    with open(os.path.join(DOCS, "HR_Task_Tracker_Oct2026.md"), "w") as f:
        f.write("\n".join(lines))

    rr = ["# HR Department - Roles & Responsibilities additions (October 2026)", "",
          "Derived from recurring instructions in the HR WhatsApp chats. To be appended to the Roles & Responsibilities document of each person so these become permanent responsibilities.", ""]
    for role, items in RR_ADDITIONS.items():
        rr.append(f"## {role}")
        rr.append("")
        for i, (title, detail) in enumerate(items, 1):
            rr.append(f"{i}. **{title}** - {detail}")
        rr.append("")
    with open(os.path.join(DOCS, "HR_Roles_Responsibilities_Additions_Oct2026.md"), "w") as f:
        f.write("\n".join(rr))

    html = ["<html><body>", "<h1>HR Department - Roles &amp; Responsibilities additions (October 2026)</h1>",
            "<p>Derived from recurring instructions in the HR WhatsApp chats (Aug 2025 - Oct 2026). To be appended to each person's Roles &amp; Responsibilities so these become permanent responsibilities. Reporting authority for both: Mukunda Charana Dasa (MKCD).</p>"]
    for role, items in RR_ADDITIONS.items():
        html.append(f"<h2>{role}</h2><ol>")
        for title, detail in items:
            html.append(f"<li><b>{title}</b> - {detail}</li>")
        html.append("</ol>")
    html.append("</body></html>")
    with open(os.path.join(OUT, "rr_additions.html"), "w") as f:
        f.write("\n".join(html))


def build_clickup(pending):
    due = datetime.strptime(DEADLINE, "%Y-%m-%d")
    # due date noon IST -> UTC 06:30 ; use date-only semantics (due_date_time False)
    due_ms = int((due - timedelta(hours=5, minutes=30)).timestamp() * 1000) + 12 * 3600 * 1000
    prio_map = {"urgent": 1, "high": 2, "normal": 3, "low": 4}
    tasks = []
    for t in pending:
        if not t["clickup"]:
            continue
        desc = (
            f"Source: {t['source']}\n"
            f"Owner as per WhatsApp: {OWNERS[t['owner']]}\n"
            f"Current status: {t['status']}\n\n"
            f"What was asked: {t['task']}\n\n"
            f"Evidence / last update: {t['evidence']}\n\n"
            f"Next action: {t['next_action']}\n\n"
            "Note: Rajendra Prasad does not use ClickUp, so this task is assigned to Gnaneshwar for tracking; "
            "coordinate with Rajendra Prasad sir where he is the owner. Deadline 10 Nov 2026. "
            "Update the HR Task Tracker Google Sheet (Master Task Log) when done."
        )
        tasks.append(dict(
            name=f"[{t['id']}] {t['task'][:110].rstrip()}" + ("..." if len(t["task"]) > 110 else ""),
            description=desc,
            priority=prio_map[t["priority"]],
            due_date=due_ms,
            due_date_time=False,
            tags=[{"name": "mkcd-followup", "tag_bg": "#1F4E78", "tag_fg": "#FFFFFF"},
                  {"name": t["category"].split(",")[0].split("&")[0].strip().lower().replace(" ", "-")[:30], "tag_bg": "#7F7F7F", "tag_fg": "#FFFFFF"}],
        ))
    with open(os.path.join(OUT, "clickup_tasks.json"), "w") as f:
        json.dump(tasks, f, indent=1)
    return tasks


if __name__ == "__main__":
    xlsx_path = os.path.join(DOCS, "HR_Task_Tracker_Oct2026.xlsx")
    pending, recurring = build_xlsx(xlsx_path)
    links = {}
    link_file = os.path.join(OUT, "links.json")
    if os.path.exists(link_file):
        links = json.load(open(link_file))
    build_markdown(pending, recurring, links.get("sheet"), links.get("clickup"), links.get("doc"))
    cu = build_clickup(pending)
    print(f"tasks={len(TASKS)} pending={len(pending)} recurring={len(recurring)} clickup={len(cu)}")
    print(f"xlsx={xlsx_path} size={os.path.getsize(xlsx_path)}")
