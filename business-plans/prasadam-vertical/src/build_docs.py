"""Builds 01_Business_Plan.docx and 04_Project_Management_Plan.docx."""
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import content as C

OUT = Path(__file__).resolve().parent.parent
MAROON = RGBColor(0x7A, 0x1F, 0x2B)
SAFFRON = RGBColor(0xC8, 0x5A, 0x00)
GREY = RGBColor(0x55, 0x55, 0x55)


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def new_doc(title, subtitle):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2))
    sec.top_margin = sec.bottom_margin = Cm(1.8)

    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for name, size, color in (("Heading 1", 18, MAROON), ("Heading 2", 13.5, SAFFRON), ("Heading 3", 11.5, MAROON)):
        h = doc.styles[name]
        h.font.name = "Calibri"
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = color

    # Cover
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Hare Krishna")
    r.font.size, r.font.color.rgb, r.bold = Pt(14), SAFFRON, True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.font.size, r.font.color.rgb, r.bold = Pt(28), MAROON, True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    r.font.size, r.font.color.rgb = Pt(15), GREY
    doc.add_paragraph()
    for line in (C.ORG, f"Prepared for: {C.PREPARED_FOR}", f"Source: {C.SOURCE_MEETING}", C.VERSION):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        r.font.size, r.font.color.rgb = Pt(11), GREY
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Figures marked [ESTIMATE] are planning assumptions. Replace them with the 6–12 month order data before committing budget.")
    r.italic, r.font.size, r.font.color.rgb = True, Pt(9), GREY
    doc.add_page_break()
    return doc


def bullets(doc, items, style="List Bullet"):
    for it in items:
        doc.add_paragraph(it, style=style)


def para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead + " ")
        r.bold = True
    p.add_run(text)
    return p


def table(doc, headers, rows, widths=None, font=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(str(h))
        r.bold, r.font.size = True, Pt(font)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, "7A1F2B")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(font)
            if ri % 2:
                shade(cells[i], "FFF4E3")
    if not widths:
        widths = [17 / len(headers)] * len(headers)
    scale = 17 / sum(widths)
    t.autofit = False
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w * scale)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w * scale)
    doc.add_paragraph()
    return t


DIAG = OUT / "diagrams"
FLOW_TITLES = [("mind_map", "Mind map overview", 17), ("order_flow", "Order-to-delivery flow (all channels)", 11),
               ("corporate_gifting_flow", "Corporate gifting sales flow", 7.5), ("festival_stall_flow", "Festival stall operations", 8.5),
               ("phase_gate_flow", "Decision and phase-gate flow", 8), ("roi_tracking_flow", "DMTT ROI tracking flow", 17)]


def diagrams(doc, slugs=None):
    for slug, title, w in FLOW_TITLES:
        if slugs and slug not in slugs:
            continue
        f = DIAG / f"{slug}.png"
        if not f.exists():
            continue
        doc.add_heading(title, 2)
        doc.add_picture(str(f), width=Cm(w))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def toc(doc, items):
    doc.add_heading("Contents", 1)
    for i, it in enumerate(items, 1):
        doc.add_paragraph(f"{i}. {it}")
    doc.add_page_break()


# ---------------------------------------------------------------------------
def business_plan():
    doc = new_doc(C.TITLE, "Detailed Business Plan")
    sections = [
        "Executive summary", "Background and meeting insights", "Vision, mission and positioning",
        "Market research", "Competitive landscape and SWOT", "Target audiences",
        "Business models and revenue streams", "New business ideas", "Products, gift tiers and packaging",
        "Pricing strategy and unit economics", "Marketing and sales plan", "Operations plan",
        "Organisation and roles", "Roadmap", "Hero campaign calendar", "Financial plan",
        "Risks and mitigation", "KPIs and ROI tracking", "Immediate action items and open questions", "Mind map and flowcharts", "Sources",
    ]
    toc(doc, sections)

    # 1
    doc.add_heading("1. Executive summary", 1)
    para(doc, C.ONE_LINE_PITCH, "The idea.")
    para(doc, "Temple visitors, donors, company owners, institutions and online followers already trust the ISKCON name. Janmashtami 2026 counters sold out at about ₹6 lakh with limited stalls. Three companies asked for Dasara gift boxes that we could not supply. A group has asked for 50,000 meals. Demand is real; what has been missing is a production owner, a sales kit and a system. That now exists: the temple production lead handles the kitchen, and Mukunda Dasa with the DMTT handles marketing, technology and sales.", "Why us, why now.")
    para(doc, "Run two price segments: a competitive mass tier for volume that keeps the kitchen busy and covers fixed costs, and a premium tier for margin. Sell through 14 revenue streams, starting with those needing little capital (internal supply, festival stalls, corporate and white-label gifting), then e-commerce, communities and catering, and later a dedicated unit, QSRs and exports.", "The model.")
    y1 = sum(r[1] for r in C.REVENUE_Y1)
    y2 = sum(r[1] * r[2] for r in C.REVENUE_Y1)
    y3 = sum(r[1] * r[2] * r[3] for r in C.REVENUE_Y1)
    ox = sum(o[1] for o in C.OPEX_Y1)
    para(doc, f"Year 1 (Oct 2026–Sep 2027) target gross revenue about ₹{y1 / 100:.1f} crore, rising to about ₹{y2 / 100:.1f} crore in Year 2 and ₹{y3 / 100:.1f} crore in Year 3 [ESTIMATE]. Year-1 incremental opex about ₹{ox} lakh and small capex under ₹30 lakh. No large building investment until demand is proven.", "The numbers.")
    para(doc, "Approve the plan in the decision meeting (target 5 Oct 2026), finish the 2-week foundation sprint (data, samples, pricing, sales kit, one hire), and use Dasara and Diwali 2026 as the pilot.", "The ask.")
    doc.add_heading("Objectives at a glance", 2)
    table(doc, ["#", "Objective", "Measure"], C.OBJECTIVES, [1.2, 6, 9.8])

    # 2
    doc.add_heading("2. Background and meeting insights", 1)
    para(doc, f"This plan is built from the {C.SOURCE_MEETING} (audio recording, transcript and summary). Every point raised is captured below, grouped by theme, so nothing is lost.")
    for theme, pts in C.MEETING_INSIGHTS:
        doc.add_heading(theme, 2)
        bullets(doc, pts)

    # 3
    doc.add_heading("3. Vision, mission and positioning", 1)
    para(doc, C.VISION, "Vision.")
    para(doc, C.MISSION, "Mission.")
    doc.add_heading("Positioning statement", 2)
    para(doc, "For families, companies and devotees who want sweets they can trust, our prasadam is the only satvik, preservative-free range prepared in a temple kitchen and offered to Sri Sri Radha Krishna, backed by the decades-old goodwill of ISKCON. Unlike local shops using colours and chemicals, every pack is pure, and every purchase supports the temple's service.")
    doc.add_heading("Brand pillars", 2)
    bullets(doc, [
        "Purity: satvik (no onion, garlic, eggs), no preservatives, no artificial colours.",
        "Devotion: prepared with care and offered to the Deity first.",
        "Trust: the ISKCON / Hare Krishna name and a long record in Hyderabad.",
        "Service: proceeds support temple activities and food-for-all programmes.",
        "Delight: attractive packaging that makes a gift feel special.",
    ])
    doc.add_heading("Guiding principles", 2)
    bullets(doc, [
        "Passive selling, not pushing: stalls where people are; people choose to buy.",
        "Start with what we have (kitchen, network, systems); invest only after demand is proven.",
        "Technology is the scale engine (like Zomato, Swiggy, Amazon): data, e-commerce, automation.",
        "Every rupee and every hour is tracked; ROI must be visible to anyone who asks.",
        "Quality and food safety are never traded for volume.",
    ])

    # 4
    doc.add_heading("4. Market research", 1)
    para(doc, "Desk research carried out on 26 September 2026. Figures are from third-party reports and should be treated as indicative.")
    table(doc, ["Topic", "Finding", "Source"], [(a, b, c) for a, b, c, _ in C.MARKET_FACTS], [3.5, 9.5, 4])
    doc.add_heading("What this means for us", 2)
    bullets(doc, [
        "The packaged-sweets market is growing about 16% a year and boxes are the leading format, which suits gift boxes.",
        "Corporate buyers drive most festival gifting; our owner-level network is the fastest route to them.",
        "The GST cut to 5% on packaged namkeen lowers consumer prices and supports the mass tier.",
        "Under the April 2026 FSSAI thresholds, Year-1 turnover below ₹1.5 crore could fall under Basic Registration, but turnover of ₹2 crore or more needs a State Licence. Plan for a State Licence from the start.",
        "Hyderabad has strong heritage brands but none with a temple and satvik story at scale, which leaves a clear gap.",
        "Other ISKCON temples (e.g. Govinda's Bangalore) prove that temple food can be sold online with same-day delivery.",
    ])

    # 5
    doc.add_heading("5. Competitive landscape and SWOT", 1)
    table(doc, ["Competitor", "Profile", "Tier", "Our edge"], C.COMPETITORS, [3.5, 5, 2, 6.5])
    doc.add_heading("SWOT analysis", 2)
    t = doc.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    fills = {"Strengths": "E8F5E9", "Weaknesses": "FFEBEE", "Opportunities": "E3F2FD", "Threats": "FFF8E1"}
    for idx, key in enumerate(["Strengths", "Weaknesses", "Opportunities", "Threats"]):
        cell = t.rows[idx // 2].cells[idx % 2]
        cell.text = ""
        r = cell.paragraphs[0].add_run(key)
        r.bold, r.font.color.rgb = True, MAROON
        for s in C.SWOT[key]:
            cell.add_paragraph("• " + s).runs[0].font.size = Pt(9)
        shade(cell, fills[key])
    doc.add_paragraph()

    # 6
    doc.add_heading("6. Target audiences", 1)
    para(doc, "Mukunda Dasa's model: different audiences, different circles and different products for each audience. Circle 1 comes to us; circle 2 is the B2B network we go to; circle 3 is the open public; circle 4 is digital; circle 5 is scale.")
    for a in C.AUDIENCES:
        doc.add_heading(f"{a['name']}  ({a['circle']})", 3)
        for k, lab in (("who", "Who"), ("needs", "Needs"), ("offer", "Our offer"), ("channel", "Channel"), ("price", "Price point")):
            para(doc, a[k], lab + ":")

    # 7
    doc.add_heading("7. Business models and revenue streams", 1)
    table(doc, ["ID", "Stream", "Phase", "Capex", "Margin"], [(s["id"], s["name"], s["phase"], s["capex"], s["margin"]) for s in C.STREAMS], [1.2, 6.5, 3, 3, 3.3])
    for s in C.STREAMS:
        doc.add_heading(f"{s['id']}. {s['name']}", 2)
        para(doc, s["what"], "What:")
        para(doc, s["why"], "Why:")
        para(doc, s["how"], "How:")
        para(doc, s["kpi"], "KPIs:")

    # 8
    doc.add_heading("8. New business ideas", 1)
    para(doc, "Additional ideas that build on the meeting and on the research. Each can be tested cheaply before scaling.")
    table(doc, ["Idea", "Description"], C.BUSINESS_IDEAS, [5, 12])

    # 9
    doc.add_heading("9. Products, gift tiers and packaging", 1)
    para(doc, "First criterion: shelf life of at least 15–20 days. Shelf-life figures are typical and must be confirmed by lab tests on our own recipes and packs.")
    table(doc, ["Product", "Category", "Shelf life (days)", "Tier", "Sweetener", "Packs", "Notes"], C.PRODUCTS, [3.2, 2.2, 1.8, 1.8, 2.2, 2.4, 3.4], font=8)
    doc.add_heading("Gift box tiers", 2)
    table(doc, ["Tier", "Price", "Contents", "Best for"], C.GIFT_TIERS, [3, 2.5, 7, 4.5])
    doc.add_heading("Packaging", 2)
    bullets(doc, C.PACKAGING)

    # 10
    doc.add_heading("10. Pricing strategy and unit economics", 1)
    para(doc, "Premium only means limited customers, limited business and no cover for capex, rent, power and salaries. Competitive only means thin margins. Run both.")
    table(doc, ["Tier", "Logic", "Used for"], C.PRICING, [3.5, 9, 4.5])
    doc.add_heading("Illustrative unit economics [ESTIMATE]", 2)
    rows = []
    for name, price, ing, pack, lab in C.UNIT_ECONOMICS:
        net = price / 1.05  # GST 5% inclusive
        gm = net - ing - pack - lab
        rows.append((name, f"₹{price}", f"₹{net:.0f}", f"₹{ing}", f"₹{pack}", f"₹{lab}", f"₹{gm:.0f}", f"{gm / net * 100:.0f}%"))
    table(doc, ["Item", "MRP", "Net of 5% GST", "Ingredients", "Packaging", "Labour & overhead", "Gross profit", "GM %"], rows, font=8.5)
    para(doc, "Replace with actual costing sheets from production (task 3.3).")

    # 11
    doc.add_heading("11. Marketing and sales plan", 1)
    table(doc, ["Lever", "Plan"], C.MARKETING, [3.5, 13.5])
    doc.add_heading("Customer funnel", 2)
    table(doc, ["Stage", "Activities"], C.MARKETING_FUNNEL, [4, 13])
    doc.add_heading("Corporate gifting sales process", 2)
    bullets(doc, [
        "List target companies from chamber, BNI, Rotary, T-Hub, donor and MNR contacts; rank by staff size.",
        "Warm introduction by Mukunda Dasa; meeting with owner/HR with sample box, catalogue and price list.",
        "Offer own-brand or white-label; share mock-up within 48 hours.",
        "Quote with volume slab, delivery date and 50% advance.",
        "Confirm PO, lock production slot, approve sticker proof.",
        "Dispatch with story card; feedback call; set reminder for next festival.",
    ], style="List Number")
    doc.add_heading("Festival stall playbook", 2)
    bullets(doc, [
        "Map footfall: entrance, darshan exit queue, Goshala walking path, staircases, garden, waterfall and benches, Goshala entrance to Samskara hall.",
        "Put stalls in a row (several counters) instead of one crowded counter; separate queues for combos.",
        "Pre-packed combos (₹100/₹200/₹500) for fast checkout; UPI QR at every counter.",
        "Runners restock hourly; stock-out alerts on WhatsApp group.",
        "Restaurant expands across the lawn with rented chairs and round tables on festival days.",
        "Leftover shelf-stable stock goes online within 48 hours with a small discount.",
    ])

    # 12
    doc.add_heading("12. Operations plan", 1)
    table(doc, ["Area", "Plan"], C.OPERATIONS, [3.5, 13.5])
    doc.add_heading("Compliance checklist", 2)
    bullets(doc, [
        "FSSAI State Licence (manufacturing) sized to expected turnover; add e-commerce and off-site stall coverage.",
        "Labels: product name, ingredients, allergens (nuts, milk, sesame), net weight, MRP, MFD, best-before, batch, FSSAI logo and number, veg mark, customer-care contact.",
        "Voluntary best-before boards on loose sweets at counters.",
        "GST registration; 5% GST on sweets and packaged namkeen (check HSN for each SKU).",
        "Legal metrology for pre-packed goods; trade licence; fire NOC; pest control; water testing; staff medical fitness.",
        "Written approval for use of the temple / ISKCON name and logo on commercial packs.",
    ])

    # 13
    doc.add_heading("13. Organisation and roles", 1)
    table(doc, ["Role", "Person", "Responsibility"], C.ORG_ROLES, [5, 4.5, 7.5])

    # 14
    doc.add_heading("14. Roadmap", 1)
    for ph in C.PHASES:
        doc.add_heading(f"{ph['name']}  ·  {ph['window']}", 2)
        para(doc, ph["goal"], "Goal:")
        bullets(doc, ph["items"])

    # 15
    doc.add_heading("15. Hero campaign calendar", 1)
    para(doc, "Dates are approximate; confirm against the temple Panchang each year.")
    table(doc, ["Festival", "When", "Opportunity", "Demand"], C.FESTIVALS, [4.5, 3.5, 6.5, 2.5])

    # 16
    doc.add_heading("16. Financial plan [ESTIMATE]", 1)
    rows, t1, t2, t3, gp1 = [], 0, 0, 0, 0
    for name, y1, m2, m3, gm, basis in C.REVENUE_Y1:
        y2, y3 = y1 * m2, y1 * m2 * m3
        t1, t2, t3, gp1 = t1 + y1, t2 + y2, t3 + y3, gp1 + y1 * gm / 100
        rows.append((name, f"{y1:.0f}", f"{y2:.0f}", f"{y3:.0f}", f"{gm}%", basis))
    rows.append(("Total", f"{t1:.0f}", f"{t2:.0f}", f"{t3:.0f}", f"{gp1 / t1 * 100:.0f}%", "Year 3 excludes QSR royalty"))
    table(doc, ["Stream", "Y1 ₹L", "Y2 ₹L", "Y3 ₹L", "GM %", "Basis"], rows, [4.5, 1.5, 1.5, 1.5, 1.4, 6.6], font=8.5)
    opex = sum(o[1] for o in C.OPEX_Y1)
    doc.add_heading("Year-1 incremental operating costs", 2)
    table(doc, ["Item", "₹ lakh", "Note"], [(a, b, c) for a, b, c in C.OPEX_Y1] + [("Total", opex, "")], [6, 2, 9])
    doc.add_heading("Year-1 summary", 2)
    bullets(doc, [
        f"Gross revenue: about ₹{t1:.0f} lakh.",
        f"Gross profit after ingredients, packaging and direct labour: about ₹{gp1:.0f} lakh ({gp1 / t1 * 100:.0f}%).",
        f"Incremental operating costs: about ₹{opex} lakh.",
        f"Net contribution to the temple: about ₹{gp1 - opex:.0f} lakh in Year 1, before capex.",
        "Break-even: the vertical should be cash-positive in the Dasara–Diwali pilot because gifting orders take 50% advance.",
    ])
    doc.add_heading("Capital expenditure (staged, only after demand is proven)", 2)
    table(doc, ["Item", "Estimate", "When"], C.CAPEX, [9, 3.5, 4.5])
    doc.add_heading("Funding", 2)
    bullets(doc, [
        "Phase 0–1: funded from existing kitchen budget plus 50% customer advances on B2B orders.",
        "Phase 2–3: reinvest vertical surplus; sponsors/donors for pagodas and machines.",
        "Phase 4+: surplus plus partner capital (QSR partner has offered funds); CSR support where appropriate.",
    ])

    # 17
    doc.add_heading("17. Risks and mitigation", 1)
    table(doc, ["ID", "Risk", "L", "I", "Score", "Mitigation", "Owner"], [(r[0], r[1], r[2], r[3], r[2] * r[3], r[4], r[5]) for r in C.RISKS], [1.3, 4.2, 0.7, 0.7, 1.1, 6.5, 2.5], font=8)

    # 18
    doc.add_heading("18. KPIs and ROI tracking", 1)
    para(doc, "Tracked by DMTT so that 'you spent this much; what did we get back?' always has a clear answer.")
    table(doc, ["KPI", "Definition", "Unit", "Frequency"], C.KPIS, [3, 8.5, 2.5, 3])
    doc.add_heading("ROI formula", 2)
    para(doc, "ROI % = (Revenue − ingredients − packaging − direct labour − marketing − allocated overheads − value of hours invested) ÷ (total costs incl. capex depreciation) × 100.")
    para(doc, "Hours are logged weekly by each team (DMTT, production, sales) and valued at a standard hourly cost, so staff time is not treated as free.")

    # 19
    doc.add_heading("19. Immediate action items and open questions", 1)
    table(doc, ["#", "Action", "Owner", "Due"], C.ACTION_ITEMS, [1.2, 8.5, 4, 3.3])
    doc.add_heading("Open questions for the decision meeting", 2)
    bullets(doc, C.OPEN_QUESTIONS)

    # 20
    doc.add_heading("20. Mind map and flowcharts", 1)
    para(doc, "The interactive, fully detailed mind map is in 03_Mind_Map.html.")
    diagrams(doc)

    # 21
    doc.add_heading("21. Sources", 1)
    for name, url in C.SOURCES:
        doc.add_paragraph(f"{name}: {url}", style="List Bullet")

    path = OUT / "01_Business_Plan.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
def pm_plan():
    doc = new_doc(C.TITLE, "Project Management Plan")
    toc(doc, [
        "Project charter", "Objectives and success criteria", "Scope", "Stakeholders", "Organisation and RACI",
        "Work breakdown structure and schedule", "Milestones", "Budget", "Risk register", "Quality management",
        "Communication plan", "Governance and reporting", "Change control", "Issue and decision log", "Action items",
    ])

    doc.add_heading("1. Project charter", 1)
    table(doc, ["Field", "Detail"], [
        ("Project name", C.TITLE + " launch"),
        ("Organisation", C.ORG),
        ("Sponsor", "Temple leadership"),
        ("Project lead", "Mukunda Dasa (marketing, sales, technology)"),
        ("Production lead", "Prabhu (temple production)"),
        ("Start / pilot end / Year-1 end", "26 Sep 2026 / 30 Nov 2026 / 30 Sep 2027"),
        ("Purpose", C.MISSION),
        ("Business case", "Proven festival demand, strong brand, warm B2B network and existing systems; two-tier pricing makes the kitchen viable; Year-1 revenue about ₹2.2 crore [ESTIMATE]."),
        ("Key constraints", "Kitchen capacity (4 staff), no historic cost data compiled, festival deadlines fixed, brand-use approval needed."),
        ("Assumptions", "Departments move all prasadam orders in-house; B2B orders take 50% advance; seasonal staff available at short notice."),
    ], [4.5, 12.5])

    doc.add_heading("2. Objectives and success criteria", 1)
    table(doc, ["#", "Objective", "Success criterion"], C.OBJECTIVES, [1.2, 6, 9.8])

    doc.add_heading("3. Scope", 1)
    doc.add_heading("In scope", 2)
    bullets(doc, C.SCOPE_IN)
    doc.add_heading("Out of scope (Year 1)", 2)
    bullets(doc, C.SCOPE_OUT)
    doc.add_heading("Deliverables", 2)
    bullets(doc, [
        "Approved business plan and charter", "Demand forecast from 6–12 months of order data",
        "Product list, recipes, costing, price list", "Sales kit: catalogue, menu, sample boxes, white-label templates",
        "Festival stall SOP and layout", "Internal weekly order system", "Shopify store and Amazon listing",
        "ROI dashboard (monthly)", "Pilot review report", "Year-1 review and Year-2 budget",
    ])

    doc.add_heading("4. Stakeholders", 1)
    table(doc, ["Stakeholder", "Interest", "Influence", "Engagement"], [
        ("Temple leadership", "Revenue, reputation, sanctity of prasadam", "High", "Quarterly steering, monthly ROI"),
        ("Production lead & kitchen team", "Workable volumes, fair staffing", "High", "Weekly plan, daily in festival weeks"),
        ("Production manager", "Resource commitments", "High", "Decision meeting, monthly review"),
        ("DMTT", "Marketing, tech, data, ROI", "High", "Weekly"),
        ("Internal departments", "Reliable supply, price", "Medium", "Order cut-offs, monthly budget"),
        ("Corporate clients & institutions", "Quality, on-time, branding", "Medium", "Account management"),
        ("Donors & devotees", "Purity, devotion, value", "Medium", "Campaigns"),
        ("QSR partner", "Brand licence, scale", "Medium", "Year-2 negotiations"),
        ("Regulators (FSSAI, GST)", "Compliance", "High", "Licences, audits"),
    ], [4, 5, 2, 6])

    doc.add_heading("5. Organisation and RACI", 1)
    table(doc, ["Role", "Person", "Responsibility"], C.ORG_ROLES, [5, 4.5, 7.5])
    doc.add_heading("RACI matrix", 2)
    para(doc, "R = Responsible, A = Accountable, C = Consulted, I = Informed.")
    table(doc, ["Activity"] + C.RACI_ROLES, C.RACI_ACTIVITIES, [4.2] + [1.6] * 8, font=8)

    doc.add_heading("6. Work breakdown structure and schedule", 1)
    para(doc, "The same WBS, with a Gantt view, is in 05_Project_Tracker.xlsx for day-to-day updates.")
    table(doc, ["ID", "Task", "Owner", "Start", "End", "Depends"], [(w[0], w[1], w[2], date.fromisoformat(w[3]).strftime("%d %b %y"), date.fromisoformat(w[4]).strftime("%d %b %y"), w[5]) for w in C.WBS], [1.1, 6.8, 3.4, 2, 2, 1.7], font=8)

    doc.add_heading("7. Milestones", 1)
    table(doc, ["Milestone", "Target date", "Gate criteria"], [
        ("M1 Plan circulated", "30 Sep 2026", "Plan shared with production lead and manager"),
        ("M2 Decision meeting / charter signed", "5 Oct 2026", "Products, quantities, budget, pricing agreed"),
        ("M3 Sales kit ready", "12 Oct 2026", "Catalogue, prices, 10 sample boxes"),
        ("M4 Dasara delivered", "20 Oct 2026", "Stalls live, first corporate orders delivered"),
        ("M5 Diwali delivered", "6 Nov 2026", "All Diwali orders dispatched on time"),
        ("M6 Pilot review", "30 Nov 2026", "P&L, ROI, lessons; go/no-go for Phase 2"),
        ("M7 Online store live", "15 Jan 2027", "Shopify live with ≥12 SKUs"),
        ("M8 Pagodas installed", "30 Jun 2027", "Permanent stalls in place"),
        ("M9 Janmashtami 2027", "Sep 2027", "Counter sales ≥ ₹30 L (base) / ₹60 L (stretch)"),
        ("M10 Year-1 review", "30 Sep 2027", "Annual report and Year-2 budget approved"),
    ], [5, 3, 9])

    doc.add_heading("8. Budget [ESTIMATE]", 1)
    opex = sum(o[1] for o in C.OPEX_Y1)
    table(doc, ["Item", "₹ lakh (Y1)", "Note"], [(a, b, c) for a, b, c in C.OPEX_Y1] + [("Total opex", opex, "")], [6, 2.5, 8.5])
    table(doc, ["Capex item", "Estimate", "When"], C.CAPEX, [9, 3.5, 4.5])

    doc.add_heading("9. Risk register", 1)
    para(doc, "Likelihood (L) and impact (I) on a 1–5 scale; score = L × I. Scores of 15+ are reviewed weekly.")
    table(doc, ["ID", "Risk", "L", "I", "Score", "Mitigation", "Owner"], [(r[0], r[1], r[2], r[3], r[2] * r[3], r[4], r[5]) for r in C.RISKS], [1.3, 4.2, 0.7, 0.7, 1.1, 6.5, 2.5], font=8)

    doc.add_heading("10. Quality management", 1)
    bullets(doc, [
        "Standard recipes and yield sheets for every SKU; taste panel sign-off before launch.",
        "Shelf-life lab test per SKU and pack type before selling online.",
        "Batch codes, MFD and best-before on every pack; FIFO in stores.",
        "Hygiene SOP: head covers, gloves, hand-wash, pest control, cleaning log; staff medical checks.",
        "Supplier checks for ghee, oil, jaggery, nuts; retain samples of each batch for 30 days.",
        "Customer complaint log with 48-hour response and root-cause review.",
        "Monthly internal audit; external FSSAI audit readiness.",
    ])

    doc.add_heading("11. Communication plan", 1)
    table(doc, ["Audience", "What", "How"], C.COMMS_PLAN, [4.5, 7, 5.5])

    doc.add_heading("12. Governance and reporting", 1)
    table(doc, ["Cadence", "Forum", "Chair"], C.GOVERNANCE, [4, 9, 4])
    doc.add_heading("KPI dashboard", 2)
    table(doc, ["KPI", "Definition", "Unit", "Frequency"], C.KPIS, [3, 8.5, 2.5, 3])
    doc.add_heading("DMTT ROI tracking", 2)
    bullets(doc, [
        "Separate cost centre for the vertical in accounts.",
        "Every order tagged by stream (R1–R14), channel and campaign.",
        "Weekly time log by each team member working on the vertical.",
        "Marketing spend tagged by campaign; ROAS per campaign.",
        "Monthly P&L by stream and a one-page ROI summary for leadership.",
    ])

    diagrams(doc, {"phase_gate_flow", "roi_tracking_flow"})

    doc.add_heading("13. Change control", 1)
    bullets(doc, [
        "Any change to scope, budget above ₹1 lakh, pricing tiers or brand use is raised as a change request.",
        "Project lead assesses impact on cost, time, capacity and risk within 3 working days.",
        "Changes up to ₹5 lakh approved by project lead and production manager; above that by the sponsor.",
        "Approved changes are logged in the tracker and communicated to all affected teams.",
    ], style="List Number")

    doc.add_heading("14. Issue and decision log", 1)
    table(doc, ["Date", "Decision / issue", "By", "Status"], [
        ("26 Sep 2026", "Two price segments (mass + premium) will run together", "Meeting", "Agreed"),
        ("26 Sep 2026", "All prasadam orders go only to the internal production lead", "Mukunda Dasa", "Agreed"),
        ("26 Sep 2026", "Shelf life of 15–20 days is the first product criterion", "Meeting", "Agreed"),
        ("26 Sep 2026", "Sugar acceptable where jaggery is too costly", "Meeting", "Agreed"),
        ("26 Sep 2026", "Weekly minimum order with fixed delivery day", "Meeting", "Agreed"),
        ("26 Sep 2026", "Start with internal demand before external expansion", "Meeting", "Agreed"),
        ("—", "Brand name to use on packs", "Sponsor", "Open"),
        ("—", "Budgets per department", "Mukunda Dasa", "Open"),
    ], [2.5, 8.5, 3, 3])

    doc.add_heading("15. Action items", 1)
    table(doc, ["#", "Action", "Owner", "Due"], C.ACTION_ITEMS, [1.2, 8.5, 4, 3.3])
    doc.add_heading("Open questions", 2)
    bullets(doc, C.OPEN_QUESTIONS)

    path = OUT / "04_Project_Management_Plan.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    print(business_plan())
    print(pm_plan())
