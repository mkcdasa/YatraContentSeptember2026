"""Builds 05_Project_Tracker.xlsx: WBS + Gantt, RACI, risks, KPIs, financial model, ROI log."""
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import content as C

OUT = Path(__file__).resolve().parent.parent / "05_Project_Tracker.xlsx"
HEAD = PatternFill("solid", fgColor="7A1F2B")
SUB = PatternFill("solid", fgColor="FBEBD0")
BAR = PatternFill("solid", fgColor="E07B00")
GROUP_BAR = PatternFill("solid", fgColor="7A1F2B")
INPUT = PatternFill("solid", fgColor="FFF8DC")
WHITE_B = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
BLUE = Font(color="1F4E9B")
thin = Side(style="thin", color="E0D6C8")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()


def sheet(title, headers, rows, widths, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font, c.alignment = HEAD, WHITE_B, Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(list(r))
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2, max_col=len(headers)):
        for c in row:
            c.alignment, c.border = WRAP, BORDER
    ws.freeze_panes = "B2"
    return ws


# 1. Read me -----------------------------------------------------------------
ws = wb.active
ws.title = "Read me"
lines = [
    (C.TITLE + " – Project Tracker", Font(bold=True, size=16, color="7A1F2B")),
    (C.ORG, None), (C.VERSION, None), ("", None),
    ("How to use", BOLD),
    ("• WBS & Gantt: update Status and % done every week; the Gantt bars are drawn by formulas from Start/End.", None),
    ("• Risk register: Score = Likelihood × Impact (formula). 15+ is red, 8–14 amber.", None),
    ("• Financial model: yellow cells are inputs (blue text). Replace [ESTIMATE] numbers with the 6–12 month order data.", None),
    ("• KPI tracker and ROI log: fill monthly/weekly; ROI is calculated automatically.", None),
    ("• Action items: the 12 actions from the 26 Sep 2026 meeting.", None),
    ("", None),
    ("Generated from src/content.py – edit there and run src/build_all.py to rebuild all documents.", Font(italic=True, color="6B5E55")),
]
for t, f in lines:
    ws.append([t])
    if f:
        ws.cell(ws.max_row, 1).font = f
ws.column_dimensions["A"].width = 120

# 2. WBS & Gantt ---------------------------------------------------------------
ws = wb.create_sheet("WBS & Gantt")
start0 = date(2026, 9, 21)  # Monday of first week
weeks = [start0 + timedelta(weeks=i) for i in range(54)]
hdr = ["ID", "Task", "Owner", "Phase", "Start", "End", "Days", "Depends on", "Deliverable", "Status", "% done"]
ws.append(hdr + [w.strftime("%d-%b") for w in weeks])
for c in ws[1]:
    c.fill, c.font = HEAD, WHITE_B
    c.alignment = Alignment(wrap_text=True, vertical="center", text_rotation=90 if c.column > len(hdr) else 0)
ws.row_dimensions[1].height = 48
for w in C.WBS:
    tid, task, owner, s, e, dep, phase, deliv = w
    r = ws.max_row + 1
    is_group = "." not in tid
    ws.append([tid, task, owner, phase, date.fromisoformat(s), date.fromisoformat(e), None, dep, deliv,
               "" if is_group else "Not started", None if is_group else 0])
    ws.cell(r, 7).value = f"=F{r}-E{r}+1"
    for col in (5, 6):
        ws.cell(r, col).number_format = "dd-mmm-yy"
    ws.cell(r, 11).number_format = "0%"
    if is_group:
        for c in ws[r][: len(hdr)]:
            c.fill, c.font = SUB, BOLD
    for i in range(len(weeks)):
        col = get_column_letter(len(hdr) + 1 + i)
        wk = weeks[i]
        d = f"DATE({wk.year},{wk.month},{wk.day})"
        ws[f"{col}{r}"] = f'=IF(AND($E{r}<={d}+6,$F{r}>={d}),1,"")'
last = ws.max_row
gantt_rng = f"{get_column_letter(len(hdr) + 1)}2:{get_column_letter(len(hdr) + len(weeks))}{last}"
ws.conditional_formatting.add(gantt_rng, FormulaRule(formula=[f'AND({get_column_letter(len(hdr) + 1)}2=1,ISERROR(FIND(".",$A2)))'], fill=GROUP_BAR, font=Font(color="7A1F2B")))
ws.conditional_formatting.add(gantt_rng, CellIsRule(operator="equal", formula=["1"], fill=BAR, font=Font(color="E07B00")))
dv = DataValidation(type="list", formula1='"Not started,In progress,Blocked,Done"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"J2:J{last}")
ws.conditional_formatting.add(f"J2:J{last}", CellIsRule(operator="equal", formula=['"Done"'], fill=PatternFill("solid", fgColor="C8E6C9")))
ws.conditional_formatting.add(f"J2:J{last}", CellIsRule(operator="equal", formula=['"Blocked"'], fill=PatternFill("solid", fgColor="FFCDD2")))
ws.conditional_formatting.add(f"J2:J{last}", CellIsRule(operator="equal", formula=['"In progress"'], fill=PatternFill("solid", fgColor="FFE0B2")))
for i, wv in enumerate([6, 44, 22, 9, 10, 10, 6, 10, 20, 12, 7], 1):
    ws.column_dimensions[get_column_letter(i)].width = wv
for i in range(len(weeks)):
    ws.column_dimensions[get_column_letter(len(hdr) + 1 + i)].width = 3.2
ws.freeze_panes = "C2"

# 3. Milestones -------------------------------------------------------------
sheet("Milestones", ["Milestone", "Target date", "Gate criteria", "Status"], [
    ("M1 Plan circulated", "30 Sep 2026", "Plan shared with production lead and manager", ""),
    ("M2 Decision meeting / charter signed", "5 Oct 2026", "Products, quantities, budget, pricing agreed", ""),
    ("M3 Sales kit ready", "12 Oct 2026", "Catalogue, prices, 10 sample boxes", ""),
    ("M4 Dasara delivered", "20 Oct 2026", "Stalls live, first corporate orders delivered", ""),
    ("M5 Diwali delivered", "6 Nov 2026", "All Diwali orders dispatched on time", ""),
    ("M6 Pilot review", "30 Nov 2026", "P&L, ROI, lessons; go/no-go for Phase 2", ""),
    ("M7 Online store live", "15 Jan 2027", "Shopify live with ≥12 SKUs", ""),
    ("M8 Pagodas installed", "30 Jun 2027", "Permanent stalls in place", ""),
    ("M9 Janmashtami 2027", "Sep 2027", "Counter sales ≥ ₹30 L (base) / ₹60 L (stretch)", ""),
    ("M10 Year-1 review", "30 Sep 2027", "Annual report and Year-2 budget approved", ""),
], [38, 14, 50, 14])

# 4. Action items ------------------------------------------------------------
sheet("Action items", ["#", "Action", "Owner", "Due", "Status", "Notes"], [a + ("Open", "") for a in C.ACTION_ITEMS], [6, 70, 30, 22, 12, 30])

# 5. RACI ----------------------------------------------------------------------
ws = sheet("RACI", ["Activity"] + C.RACI_ROLES, C.RACI_ACTIVITIES, [36] + [14] * 8)
fills = {"R": "FFE0B2", "A": "F8BBD0", "C": "E3F2FD", "I": "F5F5F5", "R/A": "FFCC80"}
for row in ws.iter_rows(min_row=2, min_col=2):
    for c in row:
        if c.value in fills:
            c.fill = PatternFill("solid", fgColor=fills[c.value])
            c.alignment = Alignment(horizontal="center")
ws.append([])
ws.append(["R = Responsible · A = Accountable · C = Consulted · I = Informed"])

# 6. Roles ------------------------------------------------------------------
sheet("Team & roles", ["Role", "Person", "Responsibility"], C.ORG_ROLES, [40, 30, 60])

# 7. Risk register -----------------------------------------------------------
ws = sheet("Risk register", ["ID", "Risk", "Likelihood (1-5)", "Impact (1-5)", "Score", "Mitigation", "Owner", "Status"],
           [(r[0], r[1], r[2], r[3], None, r[4], r[5], "Open") for r in C.RISKS], [7, 42, 11, 11, 8, 60, 20, 10])
for r in range(2, ws.max_row + 1):
    ws.cell(r, 5).value = f"=C{r}*D{r}"
    for col in (3, 4):
        ws.cell(r, col).fill, ws.cell(r, col).font = INPUT, BLUE
rng = f"E2:E{ws.max_row}"
ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["15"], fill=PatternFill("solid", fgColor="EF9A9A")))
ws.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["8", "14"], fill=PatternFill("solid", fgColor="FFE082")))
ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["8"], fill=PatternFill("solid", fgColor="C8E6C9")))

# 8. Financial model ---------------------------------------------------------
ws = wb.create_sheet("Financial model")
ws.append(["All figures ₹ lakh. Yellow cells = inputs [ESTIMATE]. Replace with real order data."])
ws["A1"].font = Font(italic=True, color="6B5E55")
ws.append(["Revenue stream", "Year 1", "Y2 growth ×", "Year 2", "Y3 growth ×", "Year 3", "Gross margin %", "Y1 gross profit", "Basis"])
for c in ws[2]:
    c.fill, c.font = HEAD, WHITE_B
first = 3
for name, y1, m2, m3, gm, basis in C.REVENUE_Y1:
    r = ws.max_row + 1
    ws.append([name, y1, m2, None, m3, None, gm / 100, None, basis])
    ws.cell(r, 4).value = f"=B{r}*C{r}"
    ws.cell(r, 6).value = f"=D{r}*E{r}"
    ws.cell(r, 8).value = f"=B{r}*G{r}"
    for col in (2, 3, 5, 7):
        ws.cell(r, col).fill, ws.cell(r, col).font = INPUT, BLUE
    ws.cell(r, 7).number_format = "0%"
lastr = ws.max_row
tot = lastr + 1
ws.append(["Total revenue", f"=SUM(B{first}:B{lastr})", "", f"=SUM(D{first}:D{lastr})", "", f"=SUM(F{first}:F{lastr})",
           f"=H{tot}/B{tot}", f"=SUM(H{first}:H{lastr})", ""])
ws.cell(tot, 7).number_format = "0%"
for c in ws[tot]:
    c.font, c.fill = BOLD, SUB
ws.append([])
ws.append(["Year-1 incremental operating costs", "₹ lakh", "Note"])
for c in ws[ws.max_row]:
    c.fill, c.font = HEAD, WHITE_B
o1 = ws.max_row + 1
for item, v, note in C.OPEX_Y1:
    ws.append([item, v, note])
    ws.cell(ws.max_row, 2).fill, ws.cell(ws.max_row, 2).font = INPUT, BLUE
o2 = ws.max_row
ws.append(["Total opex", f"=SUM(B{o1}:B{o2})"])
opex_row = ws.max_row
ws.cell(opex_row, 1).font = ws.cell(opex_row, 2).font = BOLD
ws.append([])
ws.append(["Year-1 summary", "₹ lakh"])
for c in ws[ws.max_row]:
    c.fill, c.font = HEAD, WHITE_B
ws.append(["Gross revenue", f"=B{tot}"])
ws.append(["Gross profit", f"=H{tot}"])
ws.append(["Operating costs", f"=B{opex_row}"])
ws.append(["Net contribution (before capex)", f"=H{tot}-B{opex_row}"])
ws.cell(ws.max_row, 1).font = ws.cell(ws.max_row, 2).font = BOLD
ws.append(["Net margin %", f"=B{ws.max_row}/B{tot}"])
ws.cell(ws.max_row, 2).number_format = "0.0%"
ws.append([])
ws.append(["Capex item", "Estimate", "When"])
for c in ws[ws.max_row]:
    c.fill, c.font = HEAD, WHITE_B
for row in C.CAPEX:
    ws.append(list(row))
for i, wv in enumerate([40, 12, 12, 12, 12, 12, 14, 14, 55], 1):
    ws.column_dimensions[get_column_letter(i)].width = wv
for row in ws.iter_rows(min_row=3, max_col=8):
    for c in row:
        if isinstance(c.value, (int, float)) or (isinstance(c.value, str) and c.value.startswith("=")):
            if c.number_format == "General":
                c.number_format = "#,##0.0"

# 9. Unit economics ----------------------------------------------------------
ws = sheet("Unit economics", ["Item", "MRP ₹ (incl. 5% GST)", "Net of GST", "Ingredients", "Packaging", "Labour & overhead", "Gross profit", "GM %"],
           [(n, p, None, i, pk, l, None, None) for n, p, i, pk, l in C.UNIT_ECONOMICS], [26, 16, 12, 12, 12, 16, 12, 10])
for r in range(2, ws.max_row + 1):
    ws.cell(r, 3).value = f"=B{r}/1.05"
    ws.cell(r, 7).value = f"=C{r}-D{r}-E{r}-F{r}"
    ws.cell(r, 8).value = f"=G{r}/C{r}"
    ws.cell(r, 8).number_format = "0%"
    for col in (3, 7):
        ws.cell(r, col).number_format = "#,##0"
    for col in (2, 4, 5, 6):
        ws.cell(r, col).fill, ws.cell(r, col).font = INPUT, BLUE

# 10. KPI tracker ------------------------------------------------------------
months = ["Oct-26", "Nov-26", "Dec-26", "Jan-27", "Feb-27", "Mar-27", "Apr-27", "May-27", "Jun-27", "Jul-27", "Aug-27", "Sep-27"]
ws = sheet("KPI tracker", ["KPI", "Definition", "Unit", "Frequency", "Target"] + months,
           [(k, d, u, f, "") for k, d, u, f in C.KPIS], [18, 40, 10, 11, 10] + [9] * 12)
for row in ws.iter_rows(min_row=2, min_col=5):
    for c in row:
        c.fill = INPUT

# 11. ROI & hours log -------------------------------------------------------
ws = sheet("ROI & hours log", ["Month", "Stream (R1–R14)", "Revenue ₹", "Ingredients ₹", "Packaging ₹", "Direct labour ₹", "Marketing ₹",
                                "Overheads ₹", "DMTT hours", "Production hours", "Sales hours", "Hourly cost ₹", "Value of hours ₹", "Contribution ₹", "ROI %"],
           [], [10, 16, 12, 12, 12, 13, 12, 12, 10, 12, 10, 11, 14, 14, 9])
for r in range(2, 42):
    ws.cell(r, 13).value = f'=IF(L{r}="","",(I{r}+J{r}+K{r})*L{r})'
    ws.cell(r, 14).value = f'=IF(C{r}="","",C{r}-SUM(D{r}:H{r})-N(M{r}))'
    ws.cell(r, 15).value = f'=IF(C{r}="","",N{r}/(SUM(D{r}:H{r})+N(M{r})))'
    ws.cell(r, 15).number_format = "0%"
    for col in range(1, 13):
        ws.cell(r, col).fill = INPUT
dv = DataValidation(type="list", formula1='"' + ",".join(s["id"] for s in C.STREAMS) + '"', allow_blank=True)
ws.add_data_validation(dv)
dv.add("B2:B41")

# 12. Reference sheets -------------------------------------------------------
sheet("Revenue streams", ["ID", "Stream", "What", "Why", "How", "Phase", "Capex", "Margin", "KPIs"],
      [(s["id"], s["name"], s["what"], s["why"], s["how"], s["phase"], s["capex"], s["margin"], s["kpi"]) for s in C.STREAMS],
      [6, 26, 45, 35, 45, 12, 12, 14, 35])
sheet("Audiences", ["Audience", "Circle", "Who", "Needs", "Offer", "Channel", "Price"],
      [(a["name"], a["circle"], a["who"], a["needs"], a["offer"], a["channel"], a["price"]) for a in C.AUDIENCES],
      [26, 24, 40, 30, 40, 36, 16])
sheet("Products", ["Product", "Category", "Shelf life (days)", "Tier", "Sweetener", "Pack sizes", "Notes"], C.PRODUCTS, [30, 18, 14, 14, 20, 22, 40])
sheet("Gift tiers", ["Tier", "Price", "Contents", "Best for"], C.GIFT_TIERS, [20, 18, 60, 40])
sheet("Festival calendar", ["Festival", "When (approx.)", "Opportunity", "Demand", "Forecast qty", "Staff needed", "Owner"],
      [f + ("", "", "") for f in C.FESTIVALS], [34, 22, 44, 12, 14, 14, 16])
sheet("Business ideas", ["Idea", "Description", "Priority", "Owner", "Test by"], [i + ("", "", "") for i in C.BUSINESS_IDEAS], [30, 70, 10, 16, 12])
sheet("Open questions", ["Question", "Owner", "Answer", "Date"], [(q, "", "", "") for q in C.OPEN_QUESTIONS], [70, 18, 50, 12])
sheet("Market research", ["Topic", "Finding", "Source", "URL"], C.MARKET_FACTS, [26, 70, 24, 50])

wb.save(OUT)
print(OUT)
